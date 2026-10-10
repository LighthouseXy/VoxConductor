from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
def textbox(s,x,y,w,h,text,size=18,bold=False,color='203342'):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=b.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(12)
 return b
def slide(title):
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(247,249,251);textbox(s,.5,.25,12.3,.65,title,28,True);return s
def picture(s,name,x,y,w,h):
 p=ROOT/'preview'/name;im=Image.open(p);iw,ih=im.size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;s.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
s=slide('V3-SLA｜USB方向与PCB装入修正');picture(s,'01_shell.png',.3,1,7.7,6)
textbox(s,8.2,1.4,4.6,5.5,'修复 1：接口左右镜像\nUSB、SD、麦克风声孔及线束避让同步到真实 PCB 正面。\n修复 2：底盖柱挡板\n四柱外移，壳底局部外扩；USB / SD 开口允许已焊板从底部通过。\n2026-10-11｜名义几何检查通过，待实物试装。',20)
s=slide('从USB后侧看：天线在上，USB在下');picture(s,'03_board_and_bosses.png',.2,1,7.7,6.2)
textbox(s,8.1,1.4,4.8,5.5,'USB在左，SD卡座在右。\n底盖柱中心：X=±25.5 mm\n前后Y=12 / 61.5 mm。\nPCB入口每侧预留0.5 mm；元件包络另留0.5 mm横向、0.3 mm高度。\n坐标：X=195−KiCad X，Y=KiCad Y+92，元件朝+Z。',20)
s=slide('内部装配与需要重打的零件');picture(s,'02_internal_exploded.png',.2,1,7.7,6)
textbox(s,8.1,1.3,4.8,5.6,'重打：01主壳、02底盖、04扬声器托架。\n保留：03屏幕压框、05扬声器压框、06棱镜支架。\n先将PCB固定到底盖，再将上壳垂直罩下。扬声器托架预先固定在上壳内。\n屏幕与棱镜使用原有座，排线接好后避免夹在合缝处。分解图间距仅为展示。',20)
s=slide('先试入口框，再打印整壳');textbox(s,.7,1.2,6,5.7,'第一批：trial/07入口框 + 02底盖\n入口框保留完整周圈、柱位与端口下部；用实际焊好PCB检查自由装入。\n通过后：重打01主壳、04托架\n确认USB插头能到底、SD可取出、线束不夹住，再锁紧底盖。\n不要整件打印assembly STEP；仅打印单件STL。',21)
textbox(s,7.2,1.2,5.5,5.7,'螺丝（直径×长度）\nPCB：2×5盘头自攻，3颗\n底盖：2×4沉头自攻，4颗\n托架侧面：2×8盘头自攻，2颗\n扬声器压框：2×4盘头自攻，4颗\n屏幕/棱镜组件：2×6盘头自攻，4颗\n树脂孔径与咬合未实测；手动试配，不强拧。',20)
(ROOT/'presentation').mkdir(exist_ok=True);r.save(ROOT/'presentation/V3-SLA_装配与修正说明.pptx');print('PPT saved')
